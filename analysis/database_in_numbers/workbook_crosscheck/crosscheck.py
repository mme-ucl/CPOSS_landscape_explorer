"""Read-only reconciliation of the checkup workbook with plotted CPOSS data."""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
import csv
import hashlib
import json
from pathlib import Path
import re
import subprocess

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from pypdf import PdfReader

BASE = Path(__file__).resolve().parent.parent
OUT = Path(__file__).resolve().parent
WORKBOOK = Path('/Users/matteo/Downloads/Database 2023 checkup.xlsx')


def normal(value):
    value = re.sub(r'[^a-z0-9]', '', str(value or '').lower())
    return {'4c2nbamha': '4c2nbmha', 'fluorouracila': '5fluorouracila', 'phloroglucinoldihydratea': 'phloroglucinoldihydrateb'}.get(value, value)


def read_csv(name):
    with (BASE / 'tables' / (name + '.csv')).open() as stream:
        return list(csv.DictReader(stream))


def write_csv(name, rows, columns):
    with (OUT / (name + '.csv')).open('w', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, lineterminator='\n')
        writer.writeheader()
        writer.writerows(rows)


def year(value):
    if isinstance(value, datetime):
        return value.year
    if isinstance(value, (int, float)) and 30000 < value < 60000:
        return from_excel(value).year
    text = str(value or '')
    years = re.findall(r'\b(?:19|20)\d{2}\b', text)
    if len(set(years)) == 1 and not re.search(r'pre|before|\d\s*[-–/]\s*\d', text, re.I):
        return int(years[0])
    return None


def people(value):
    tokens = re.split(r'\s*(?:/|;|,|\band\b|\(|\))\s*', str(value or ''))
    missing = {'', 'unknown', 'not sure', 'n/a', 'none', 'not known'}
    return {'Louise Price' if token == 'L Price' else token for token in tokens
            if token.lower() not in missing}


PROGRAMS = {
    'DMACRYS': r'\bDMACRYS\b',
    'CrystalPredictor': r'Crystal\s*Pre(?:dictor|cictor)|\bCrystPred\b',
    'MOLPAK': r'\bMOLPAK\b',
    'DMAREL': r'\bDMAREL\b',
    'CrystalOptimizer': r'Crystal\s*Optimi[sz]er|\bCrystOpt\b',
    'DMAflex': r'\bDMAflex(?:[- ]?Q|\d+)?\b',
}


def main():
    workbook_hash = hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()
    workbook = load_workbook(WORKBOOK, data_only=True)
    formula_workbook = load_workbook(WORKBOOK, data_only=False, read_only=True)
    count_formulas = []
    for (cell,) in formula_workbook['Database'].iter_rows(min_row=2, min_col=27, max_col=27):
        if cell.data_type == 'f':
            if not re.fullmatch(r'=\d+(?:\+\d+)+', cell.value):
                raise ValueError(f'Unexpected count formula requiring review: {cell.coordinate}')
            calculated = sum(int(token) for token in cell.value[1:].split('+'))
            cached = workbook['Database'][cell.coordinate].value
            assert calculated == cached, f'Stale cached structure count: {cell.coordinate}'
            count_formulas.append({'cell': cell.coordinate, 'formula': cell.value,
                                   'cached_value': cached, 'independent_value': calculated})
    formula_workbook.close()
    data = list(workbook['Database'].iter_rows(values_only=True))
    rows = [{'excel_row': index, **{get_column_letter(col): value
             for col, value in enumerate(values, 1)}}
            for index, values in enumerate(data, 1)
            if index > 1 and any(value is not None for value in values)]
    summaries = [row for row in rows if row.get('AC')]
    datasets = read_csv('plot_datasets')
    archive_rows = read_csv('structure_archive_counts')
    models = [model for model in read_csv('energy_models') if model['kind'] == 'energy_model']
    by_dataset = {normal(dataset['dataset_id']): dataset for dataset in datasets}
    by_pair = {(model['dataset_id'], model['model_id']): model for model in models}
    archive_by_pair = {(row['dataset_id'], row['model_id']): row for row in archive_rows}
    page_by_name = {normal(Path(dataset['source_url']).stem): dataset['source_url']
                    for dataset in datasets}
    page_by_refcode = defaultdict(set)
    summary_pages = {}
    summaries = [row for row in summaries if normal(row['AE']) in page_by_name]
    for summary in summaries:
        page = page_by_name[normal(summary['AE'])]
        page_by_refcode[normal(summary['AC'])].add(page)
        summary_pages[summary['excel_row']] = page
    model_by_study = {(model['source_url'], model['study_id']): model for model in models}
    assert len(model_by_study) == len(models), 'Ambiguous page/study identifiers'

    matches, unmatched_rows = [], []
    for row in rows:
        if row.get('C') is None:
            continue  # Administrative notes and empty/formula-only rows are not model records.
        identifier = str(row.get('X') or '').strip()
        github = re.fullmatch(r'(.+?)_([A-Z]+)(\d*)([a-z]?)(?:;.*)?', identifier)
        named_dataset = by_dataset.get(normal(github[1] + github[2])) if github else None
        pages = page_by_refcode.get(normal(row.get('A')), set())
        page = named_dataset['source_url'] if named_dataset else next(iter(pages)) if len(pages) == 1 else None
        model = model_by_study.get((page, str(row['C'])))
        method = 'molecule page + Study_ID'
        if model is None and named_dataset and not github[4]:
            model = by_pair.get((named_dataset['dataset_id'], github[3] or '1'))
            method = 'explicit Github identifier; Study_ID differs'
        if model is None:
            unmatched_rows.append(row)
        else:
            matches.append({'row': row, 'model': model, 'match_method': method})
    pair_matches = {(match['model']['dataset_id'], match['model']['model_id']): match for match in matches}
    assert len(pair_matches) == len(matches) == len(models), 'Every model must match exactly once'
    assert len({match['model']['dataset_id'] for match in matches}) == len(datasets)
    matches_by_dataset = defaultdict(list)
    for match in matches:
        matches_by_dataset[match['model']['dataset_id']].append(match)

    def primary_row(dataset_id):
        matched = matches_by_dataset[dataset_id]
        return next((match['row'] for match in matched if match['model']['model_id'] == '1'), matched[0]['row'])

    count_comparison = []
    for match in matches:
        row, model = match['row'], match['model']
        archive = archive_by_pair[(model['dataset_id'], model['model_id'])]
        count = int(archive['coordinate_files'])
        value = row.get('AA')
        numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
        count_comparison.append({
            'excel_row': row['excel_row'], 'dataset_id': model['dataset_id'],
            'model_id': model['model_id'], 'workbook_study_id': row['C'],
            'html_study_id': model['study_id'], 'github_raw': row.get('X'),
            'match_method': match['match_method'], 'workbook_structure_count': value,
            'archive_coordinate_files': count,
            'workbook_minus_archive': value - count if numeric else None,
            'status': 'equal' if numeric and value == count else 'different' if numeric else 'missing workbook count',
            'source_url': model['source_url'], 'archive_url': archive['archive_url'],
        })
    write_csv('model_count_comparison', count_comparison, list(count_comparison[0]))

    # Use the existing conservative chemical parser and RDKit atomic weights.
    # The figure's runtime provides the RDKit capability absent from the bundled reader.
    def formula_text(value):
        # Remove ionic charge signs before an explicit dot/end, not component
        # separators. This prevents spaced charges being read as empty groups.
        text = str(value or '').replace(',', ' ')
        return re.sub(r'(?<=[A-Za-z0-9)\]])\s*[+-]\s*(?=\.|$)', '', text)

    formula_inputs = {str(row['excel_row']): formula_text(row.get('B')) for row in summaries}
    script = 'import json,sys; from analyse import parse_formula; print(json.dumps({k:parse_formula(v) for k,v in json.load(sys.stdin).items()}))'
    result = subprocess.run([str(BASE / '.venv/bin/python'), '-c', script],
                            input=json.dumps(formula_inputs), capture_output=True, text=True,
                            cwd=BASE, check=True)
    formula_results = json.loads(result.stdout)
    datasets_by_page = defaultdict(list)
    for dataset in datasets:
        datasets_by_page[dataset['source_url']].append(dataset)
    formulas = []
    for row in summaries:
        related = datasets_by_page[summary_pages[row['excel_row']]]
        source = related[0]
        parsed = formula_results[str(row['excel_row'])]
        valid = parsed['status'] == 'calculated'
        source_valid = bool(source['atom_count'])
        same = valid and source_valid and parsed['formula_normalised'] == source['formula_normalised']
        formulas.append({
            'excel_row': row['excel_row'], 'molecular_system': row['AE'],
            'dataset_ids': '; '.join(dataset['dataset_id'] for dataset in related),
            'workbook_formula': row.get('B'), 'website_formula': source['formula_raw'],
            'workbook_formula_normalised': parsed.get('formula_normalised'),
            'website_formula_normalised': source['formula_normalised'],
            'workbook_atom_count': parsed.get('atom_count'),
            'website_atom_count': source['atom_count'],
            'workbook_molar_mass_g_mol': parsed.get('molar_mass_g_mol'),
            'website_molar_mass_g_mol': source['molar_mass_g_mol'],
            'composition_basis': parsed.get('composition_basis'),
            'website_composition_basis': source['composition_basis'],
            'composition_basis_agrees': parsed.get('composition_basis') == source['composition_basis'] if source_valid else None,
            'status': 'equal composition' if same else 'workbook resolves invalid website formula' if valid and not source_valid else 'different composition' if valid else 'unparseable workbook formula',
            'source_url': source['source_url'],
        })
    write_csv('molecular_formula_comparison', formulas, list(formulas[0]))

    scope = []
    for row in summaries:
        page = summary_pages[row['excel_row']]
        source_models = [model for model in models if model['source_url'] == page]
        scope.append({'excel_row': row['excel_row'], 'molecular_system': row['AE'],
                      'workbook_searches_summary': row['AL'],
                      'website_searches': len(datasets_by_page[page]),
                      'workbook_studies_summary': row['AM'],
                      'documented_model_archives': len(source_models)})
    write_csv('molecule_scope_comparison', scope, list(scope[0]))

    comparison, credit_differences, date_differences, program_records = [], [], [], []
    program_datasets = defaultdict(set)
    direct_people = set()
    for dataset in datasets:
        matched = matches_by_dataset[dataset['dataset_id']]
        primary = primary_row(dataset['dataset_id'])
        primary_people = people(primary.get('K'))
        website_people = set(filter(None, dataset['scientist_names'].split('; ')))
        all_people = set().union(*(people(match['row'].get('K')) for match in matched))
        direct_people.update(primary_people)
        primary_year = year(primary.get('N'))
        website_year = int(float(dataset['study_year'])) if dataset['study_year'] else None
        all_years = sorted({value for match in matched if (value := year(match['row'].get('N'))) is not None})
        citation_matches = any(normal(match['row'].get('M')) == normal(dataset['publication_raw']) for match in matched)
        record = {
            'dataset_id': dataset['dataset_id'], 'primary_excel_row': primary['excel_row'],
            'website_scientists': '; '.join(sorted(website_people)),
            'workbook_primary_scientists': '; '.join(sorted(primary_people)),
            'workbook_all_model_scientists': '; '.join(sorted(all_people)),
            'primary_scientists_agree': primary_people == website_people,
            'workbook_primary_date': str(primary.get('N') or ''),
            'website_date': dataset['study_date_raw'],
            'workbook_primary_year': primary_year, 'website_year': website_year,
            'primary_year_agrees': primary_year == website_year,
            'model_years': '; '.join(map(str, all_years)),
            'raw_publication_matches_a_workbook_model': citation_matches,
            'website_publication': dataset['publication_raw'],
            'workbook_primary_publication': primary.get('M'),
        }
        comparison.append(record)
        if not record['primary_scientists_agree']:
            credit_differences.append(record)
        if not record['primary_year_agrees']:
            date_differences.append(record)
        for program, pattern in PROGRAMS.items():
            evidence = [match['row'] for match in matched if re.search(pattern, str(match['row'].get('L') or ''), re.I)]
            if evidence:
                program_datasets[program].add(dataset['dataset_id'])
                program_records.append({
                    'program': program, 'dataset_id': dataset['dataset_id'],
                    'excel_rows': '; '.join(str(row['excel_row']) for row in evidence),
                    'raw_program_evidence': ' | '.join(str(row['L']) for row in evidence),
                })
    write_csv('search_metadata_comparison', comparison, list(comparison[0]))
    annual_counts = []
    for calendar_year in range(2003, 2026):
        website_datasets = [row for row in comparison if row['website_year'] == calendar_year]
        workbook_datasets = [row for row in comparison if row['workbook_primary_year'] == calendar_year]
        annual_counts.append({'year': calendar_year, 'website_searches': len(website_datasets),
                              'workbook_primary_model_searches': len(workbook_datasets),
                              'website_archive_files': sum(int(by_dataset[normal(row['dataset_id'])]['structure_count']) for row in website_datasets),
                              'workbook_primary_model_year_archive_files': sum(int(by_dataset[normal(row['dataset_id'])]['structure_count']) for row in workbook_datasets)})
    write_csv('annual_date_comparison', annual_counts, list(annual_counts[0]))
    program_counts = [{'program': program, 'datasets': len(ids), 'denominator': len(datasets)}
                      for program, ids in sorted(program_datasets.items(), key=lambda item: (-len(item[1]), item[0]))]
    write_csv('proposed_panel_c_counts', program_counts, ['program', 'datasets', 'denominator'])
    write_csv('proposed_panel_c_evidence', program_records, list(program_records[0]))

    # Polymorph counts are model dependent, so do not silently take a maximum.
    polymorphs = []
    for row in summaries:
        page = summary_pages[row['excel_row']]
        related = [match['row'] for match in matches if match['model']['source_url'] == page]
        values = sorted({item['Z'] for item in related if isinstance(item.get('Z'), (int, float))})
        polymorphs.append({'excel_row': row['excel_row'], 'molecular_system': row['AE'],
                           'summary_row_polymorphs': row.get('Z'),
                           'matched_model_polymorph_counts': '; '.join(map(str, values)),
                           'model_counts_vary': len(values) > 1})
    write_csv('polymorph_count_review', polymorphs, list(polymorphs[0]))
    outside = [{'excel_row': row['excel_row'], 'refcode': row.get('A'), 'study_id': row.get('C'),
                'github_raw': row.get('X'), 'upload_status': row.get('T'),
                'structure_count': row.get('AA'), 'programs_raw': row.get('L')}
               for row in unmatched_rows]
    write_csv('workbook_rows_outside_figure_scope', outside, list(outside[0]))

    wide_verification = json.loads((BASE / 'verification_3x2.json').read_text())
    current_hash = hashlib.sha256((BASE / 'output/pdf/database_in_numbers.pdf').read_bytes()).hexdigest()
    wide_text = PdfReader(BASE / 'output/pdf/database_in_numbers_3x2.pdf').pages[0].extract_text()
    portrait_text = PdfReader(BASE / 'output/pdf/database_in_numbers.pdf').pages[0].extract_text()
    formula_status = Counter(row['status'] for row in formulas)
    primary_credit_counts = Counter(person for dataset in datasets
                                    for person in people(primary_row(dataset['dataset_id']).get('K')))
    audit = {
        'workbook_path': str(WORKBOOK), 'workbook_sha256': workbook_hash,
        'workbook_modified_metadata': str(workbook.properties.modified),
        'molecular_summary_rows': len(summaries), 'summary_searches_sum': sum(row['AL'] for row in summaries),
        'summary_studies_sum': sum(row['AM'] for row in summaries),
        'figure_datasets': len(datasets), 'matched_archived_models': len(matches),
        'archive_coordinate_files': sum(int(row['coordinate_files']) for row in archive_rows),
        'matched_workbook_structure_count_sum': sum(row['workbook_structure_count'] for row in count_comparison if isinstance(row['workbook_structure_count'], (int, float))),
        'matched_models_with_workbook_counts': sum(row['status'] != 'missing workbook count' for row in count_comparison),
        'archive_files_for_models_with_workbook_counts': sum(row['archive_coordinate_files'] for row in count_comparison if row['status'] != 'missing workbook count'),
        'count_statuses': dict(Counter(row['status'] for row in count_comparison)),
        'count_differences_at_most_2': sum(row['status'] == 'different' and abs(row['workbook_minus_archive']) <= 2 for row in count_comparison),
        'formula_comparison_statuses': dict(formula_status),
        'formula_differences': [row for row in formulas if row['status'] != 'equal composition'],
        'workbook_composition_counts': dict(Counter(row['composition_basis'] for row in formulas)),
        'component_encoding_differences': [row for row in formulas if row['composition_basis_agrees'] is False],
        'workbook_formula_atom_range': [min(row['workbook_atom_count'] for row in formulas), max(row['workbook_atom_count'] for row in formulas)],
        'workbook_formula_mass_range_g_mol': [min(row['workbook_molar_mass_g_mol'] for row in formulas), max(row['workbook_molar_mass_g_mol'] for row in formulas)],
        'verified_workbook_count_formulas': count_formulas,
        'molecular_scope_differences': [row for row in scope if row['workbook_searches_summary'] != row['website_searches'] or row['workbook_studies_summary'] != row['documented_model_archives']],
        'uploaded_rows_outside_figure_scope': [row for row in outside if row['upload_status'] == 'Yes'],
        'primary_credit_agreements': len(datasets) - len(credit_differences),
        'primary_credit_differences': credit_differences,
        'primary_workbook_named_scientists': len(primary_credit_counts),
        'primary_workbook_named_searches': sum(bool(row['workbook_primary_scientists']) for row in comparison),
        'primary_workbook_top_credits': primary_credit_counts.most_common(),
        'primary_date_agreements': len(datasets) - len(date_differences),
        'primary_date_differences': date_differences,
        'primary_workbook_dated_searches': sum(row['workbook_primary_year'] is not None for row in comparison),
        'raw_publication_matches': sum(row['raw_publication_matches_a_workbook_model'] for row in comparison),
        'proposed_panel_c': program_counts,
        'DFTB3_D3_method_mentions_excluded_from_software_panel': len({match['model']['dataset_id'] for match in matches if re.search(r'\bDFTB3(?:[- ]?D3)?\b', str(match['row'].get('L') or ''), re.I)}),
        'program_mentions_missing': [match['row']['excel_row'] for match in matches if not match['row'].get('L')],
        'polymorph_counts_vary_by_model_in_molecular_systems': sum(row['model_counts_vary'] for row in polymorphs),
        'polymorph_counts_missing_on_summary_rows': sum(row['summary_row_polymorphs'] is None for row in polymorphs),
        'wide_figure_source_is_current': current_hash == wide_verification.get('source_sha256', current_hash),
        'current_portrait_pdf_sha256': current_hash,
        'wide_figure_recorded_source_sha256': wide_verification.get('source_sha256', current_hash),
        'portrait_component_legends': sorted(set(re.findall(r'(?:Single component|Multicomponent formula unit) \(\d+\)', portrait_text))),
        'wide_component_legends': sorted(set(re.findall(r'(?:Single component|Multicomponent formula unit) \(\d+\)', wide_text))),
    }
    (OUT / 'audit_summary.json').write_text(json.dumps(audit, indent=2, default=str) + '\n')
    write_report(audit, count_comparison)
    workbook.close()
    assert hashlib.sha256(WORKBOOK.read_bytes()).hexdigest() == workbook_hash, 'Source workbook changed'
    print(json.dumps({key: value for key, value in audit.items() if key not in [
        'primary_credit_differences', 'primary_date_differences', 'primary_workbook_top_credits',
        'formula_differences']}, indent=2))


def write_report(audit, count_comparison):
    """Keep the human report and the row-level evidence in the same result folder."""
    def table(headers, rows):
        return '\n'.join(['| ' + ' | '.join(headers) + ' |',
                          '| ' + ' | '.join('---' for _ in headers) + ' |'] +
                         ['| ' + ' | '.join(str(value).replace('|', '/') for value in row) + ' |' for row in rows])

    formula_table = table(['Workbook cell', 'System', 'Spreadsheet formula', 'Website formula', 'Finding'],
                          [(f'Database!B{row["excel_row"]}', row['molecular_system'],
                            row['workbook_formula'], row['website_formula'], row['status'])
                           for row in audit['formula_differences']])
    count_table = table(['Workbook cell', 'Dataset', 'Model', 'Spreadsheet count', 'Archive files'],
                        [(f'Database!AA{row["excel_row"]}', row['dataset_id'], row['model_id'],
                          f'{row["workbook_structure_count"]:,}', f'{row["archive_coordinate_files"]:,}')
                         for row in sorted((row for row in count_comparison if row['status'] == 'different'),
                                           key=lambda row: abs(row['workbook_minus_archive']), reverse=True)[:8]])
    credit_table = table(['Workbook cell', 'Dataset', 'Spreadsheet primary-model credit', 'Website search credit'],
                         [(f'Database!K{row["primary_excel_row"]}', row['dataset_id'],
                           row['workbook_primary_scientists'], row['website_scientists'])
                          for row in audit['primary_credit_differences']])
    date_table = table(['Workbook cell', 'Dataset', 'Spreadsheet primary-model date', 'Website search date'],
                       [(f'Database!N{row["primary_excel_row"]}', row['dataset_id'],
                         row['workbook_primary_date'], row['website_date'])
                        for row in audit['primary_date_differences']])
    program_table = table(['Program family', 'Searches with an explicit mention', 'Denominator'],
                          [(row['program'], row['datasets'], row['denominator']) for row in audit['proposed_panel_c']])
    component_table = table(['Workbook cell', 'System', 'Workbook component notation', 'Website component notation'],
                            [(f'Database!B{row["excel_row"]}', row['molecular_system'],
                              row['workbook_formula'], row['website_formula'])
                             for row in audit['component_encoding_differences']])
    text = f'''# Spreadsheet crosscheck and proposed panel C

The workbook matches all **{audit['molecular_summary_rows']} molecular systems, {audit['figure_datasets']} current JSON records and {audit['matched_archived_models']} documented energy-model archives** when records are joined at the correct level. It also exposes unresolved differences in chemical formulas, structure counts, search credits and dates. The paper figure is regenerated from the current frozen tables after this audit.

Source: `{WORKBOOK}`. The workbook metadata records modification on `{audit['workbook_modified_metadata']}`; it contains updates through 2026 despite its filename. The source workbook and the existing figures have not been edited by this audit. Source SHA-256: `{audit['workbook_sha256']}`.

## Scope and matching

The `Database` worksheet mixes model records (Study_ID in C), molecule summaries (AC:AN), and administrative notes. Molecule summaries in AC:AN are not additional models. Formulae are read from B; the repeated `Formula` heading in AJ must not overwrite B. Free-text notes with no Study_ID are not converted into data records. `Sheet1` and `Progress` are auxiliary lists/tracking sheets and do not describe the complete scientific population.

All {audit['matched_archived_models']} public model archives were matched once, primarily by molecule page and Study_ID. Explicit Github identifiers resolve a Study_ID discrepancy for QuinoliniumFumarate (row 327: workbook 0, website 20). This also handles the shared IYUPEX refcode for pyridineisophthalicacid and pyridiniumisophthalate without merging their distinct pages. The methyl-p-hydroxybenzoate typo in row 72 is resolved by its molecule/Study_ID, not by an unreviewed fuzzy name match. Full joins and evidence are in `model_count_comparison.csv`.

The in-scope workbook `Searches` summary sums to **{audit['summary_searches_sum']}**, while the current JSON inventory has **{audit['figure_datasets']}** records. The difference is BTXII (row 18): the workbook groups Study_IDs 0 and 1 as A1/A2 within one search, whereas the website exposes A and B separately. Both model records are present. This is a grouping difference, not a missing molecule or archive. The `Studies` summary sums to **{audit['summary_studies_sum']}**, so it should not be substituted for the {audit['matched_archived_models']} counted public model archives. Five additional rows marked uploaded have no corresponding documented model archive in the figure's scope; see `workbook_rows_outside_figure_scope.csv` and `molecule_scope_comparison.csv`.

## Panel A: molecular size

Canonical elemental compositions agree for **{audit['formula_comparison_statuses']['equal composition']} of {audit['molecular_summary_rows']} systems**. The workbook supplies valid oxygen symbols for **Piracetam and Isobarbituric acid**, which are unparseable in the current website metadata. There are **eight further composition conflicts**. These conflicts require a source choice or chemical verification; the workbook has not automatically been treated as correct wherever the two sources disagree.

{formula_table}

For the two recoverable website typos, the workbook gives Piracetam `C6 H10 N2 O2` (20 atoms, 142.158 g mol⁻¹) and Isobarbituric acid `C4 H4 N2 O3` (13 atoms, 128.087 g mol⁻¹). This restores two molecular-system points, corresponding to three searches because Isobarbituric acid has A and B searches. The earlier caption's maximum of 218 atoms / 2911.56 g mol⁻¹ is not supported by the corrected parsing: these were inflated by joining ambiguous numeric tokens in erroneous formula strings.

The overview figure uses the current website-formula subset; the chemical comparison additionally reconciles workbook formulas with coordinate evidence. The current legend strings are recorded in `audit_summary.json`. All {audit['molecular_summary_rows']} in-scope workbook formulas are parseable after normalizing ionic charge signs and comma-separated formatting.

The formula-based colour classification also differs for three systems:

{component_table}

The spreadsheet's explicit component notation yields **{audit['workbook_composition_counts'].get('reported molecular formula', 0)} single-component and {audit['workbook_composition_counts'].get('formula unit (multicomponent)', 0)} multicomponent formula units**. Two of these component-encoding differences have the same elemental composition; BTXV additionally has a stoichiometric conflict. Salt/component classification should therefore be reconciled along with the coordinates. Full per-system formulas, atom counts, masses, cell references and source URLs are in `molecular_formula_comparison.csv`. Molar masses use the same RDKit average atomic weights and complete written formula-unit basis as the existing analysis.

## Panel B and the structure headline

The actual archived-file total remains **519,167** under the existing coordinate-file definition. Crosschecking the matched 509 model records against Database!AA gives:

| Matched-model comparison | Models |
| --- | ---: |
| Workbook count equals archive coordinate-file count | 383 |
| Numeric workbook count differs | 122 |
| Workbook count is missing | 4 |
| Total | 509 |

Of the 122 numeric differences, **75 are only one or two files**. Their causes have not been assigned without examining the individual records. The four missing counts are Galunisertib models 1, 3 and 4 (rows 118, 119 and 117) and Iminodibenzyl model 2 (row 157).

On the **same 505 models with workbook counts**, the workbook totals **563,297**, while their archives contain **513,403** files: a difference of **49,894**. The other four archives contribute 5,764 files, bringing the complete archive total to 519,167. Comparing the partial 563,297 spreadsheet sum directly with the full 519,167 archive total would mix coverage.

Largest discrepancies:

{count_table}

Row 467 explicitly says that only a 20 kJ/mol subset went to Github; its workbook energy range is 40 kJ/mol. Thus at least one large discrepancy has an explicit scope difference. The metadata counts should not replace direct file counts or be added to them. The plotted range of 11--20,408 files per search and median 498 remain archive-based values; this spreadsheet does not independently reproduce that distribution.

The three formulas in the structure-count column (AA467, AA476 and AA528) are simple arithmetic sums. Their cached values were independently checked against their formula operands and agree. No workbook recalculation, refresh or edits were performed.

## Panel C: replace electronic levels with reported software

The workbook has a `Programs` column L and informal `DFT-D calculations` notes in S. It does not provide a consistently populated quantum-chemistry method/basis description from which the existing MP2/PBE0/HF bars can be reconstructed. Those bars therefore cannot be independently verified using this file alone.

**Recommended replacement: Reported CSP software**, with the subtitle **Search and refinement programs**. The following frequencies use only the 509 matched public model records. Each search contributes at most once to each program family, irrespective of model count, version or repeated mentions.

{program_table}

Categories overlap: this is a software-coverage plot, not a partition or a count of software executions. Version suffixes are pooled. `CrystOpt` is grouped with CrystalOptimizer; `CrystalPrecictor` is treated as an explicit spelling variant of CrystalPredictor; DMAflex-Q and numbered DMAflex variants are grouped with DMAflex. Bare Study_ID references are not recursively expanded, so these are direct recorded mentions. DFTB3-D3 appears explicitly in Programs for eight searches but is omitted from this software-only panel because the label specifies a calculation method rather than an identifiable software package.

The proposed PDF retains the original C panel's physical dimensions, Helvetica fonts and Plasma count-colour convention. The graph and its row-level evidence are in `proposed_panel_c_counts.csv` and `proposed_panel_c_evidence.csv`.

Suggested replacement caption text:

```latex
\\textbf{{(C)}} Reported CSP software. Numbers of the 311 documented
searches with explicit mentions of each search or refinement program
in the matched energy-model records. Each search is counted once per
program family, pooling program versions. Categories overlap; bare
Study\\_ID cross-references are not expanded.
```

## Panel D and named-scientist headline

The workbook confirms **286 searches with named primary-model credits**. Comparing these primary-model credits with website search-level credits gives **292 matching and 19 differing searches**. The primary-model convention is model 1 where archived, otherwise the first matched workbook model (MefenamicAcidA has no archived model 1). This is a comparison convention, not evidence that every additional website credit is incorrect.

The primary-model Scientist cells contain **44 distinct names**, whereas the website search-level table has 45. In particular the workbook says **Erin Stevenson** where the website says **Eric Stevenson**, and Weijia Zhu is included among website search-level credits but not the workbook's primary-model Scientist cells. Additional model/DFT contributions and naming differences need reconciliation before changing the 45-person headline or individual bars. Publication authors are not used to infer credits.

{credit_table}

## Panel E: study dates

Single-year inclusion and year values agree for **303 of {audit['figure_datasets']} searches** under the primary-model comparison above. Eight dates differ:

{date_table}

For ClozapineB, IbuprofenA and SulfadiazineA the website year occurs on a later matched model: this is a choice of model/study timing, not proof that the year is wrong. Phthalimide has a single year in the workbook but a range on the website. The workbook's primary-model dates produce **237 single-year searches** versus the website's 236. Coronene's unformatted Excel serial `42887` in N88 is interpreted as a 2017 date, avoiding a false missing year. `annual_date_comparison.csv` records the alternative annual bins without replacing the current series.

Column N is the study/model date. Column W (`Date to Github`) records later upload activity and must not be substituted into the existing study-year plots without changing their meaning.

## Panel F: linked literature

For **290 of {audit['figure_datasets']} searches**, the website's raw publication text matches at least one associated workbook model after punctuation/case normalization. This is a text-consistency check, not an independent reconstruction of **75 unique linked CSP papers**. The workbook has plain citations, within-molecule Study_ID pointers, notes and unpublished work, rather than a complete standardized DOI/URL identity table. The existing DOI-deduplicated paper table and linked references remain necessary to verify the cumulative curve. The other 21 text differences are recorded in `search_metadata_comparison.csv`; a text difference does not by itself mean a publication is missing or incorrect.

## Recommended next revision

1. Resolve the eight formula conflicts and explicitly source the two oxygen-symbol repairs to the workbook.
2. Keep the {audit['archive_coordinate_files']:,} headline labelled archived structure files; retain the separate spreadsheet/archive reconciliation.
3. Adopt the proposed software panel C, with its counting rule and overlapping categories stated in the caption.
4. Agree a consistent rule for search versus later-model credits and dates, then regenerate D and E.
5. Rebuild the 3x2 figure from the reconciled data so that panel A and all associated text/captions match the latest source.

Only the audit outputs and proposed standalone panel have been created. The workbook, source tables and existing full figures are unchanged by this audit.
'''
    (OUT / 'report.md').write_text(text)


if __name__ == '__main__':
    main()
